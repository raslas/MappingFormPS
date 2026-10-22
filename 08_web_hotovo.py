# -*- coding: utf-8 -*-
"""
Naplní webové SQLite databázy pre prezeračku hotových území
(C:\\wamp64\\www\\MapovaniePrePSHotovo).

Zdroj údajov:  <cloud>\\SKUEV####\\MapovaniePrePS.gpkg (+ AktivityLookup.gpkg
               v tom istom priečinku) – teda tie isté súbory, z ktorých
               exportujú 05_export_shapefile.py a 06_naplnenie_access.py,
               a fotky z <cloud>\\SKUEV####\\f (+ SKUEV####_FotoTable_*.txt).
Výstup:        C:\\wamp64\\www\\MapovaniePrePSHotovo\\data\\SKUEV####.sqlite
                   (jedna databáza na územie – na server sa nahrá len
                   zmenené územie)
               C:\\wamp64\\www\\MapovaniePrePSHotovo\\data\\foto\\SKUEV####\\
                   zmenšené fotky (+ podpriečinok "n" s náhľadmi)

Na web sa dostane len to, čo sem pošleš argumentom – "hotové" územie je
to, ktoré si sem pridal:

    python 08_web_hotovo.py SKUEV0870
    python 08_web_hotovo.py SKUEV0870 SKUEV0862 0817
    python 08_web_hotovo.py --remove SKUEV0870      # vyhodí územie z webu
    python 08_web_hotovo.py --list                  # čo je práve na webe
    python 08_web_hotovo.py SKUEV0870 --bez-fotiek  # rýchly refresh bez fotiek

Opakované spustenie pre to isté územie jeho databázu vytvorí nanovo (zostaví sa
pod dočasným názvom a až potom nahradí starú), takže sa dá po každej zmene
v cloude spustiť znova. Na server stačí nahrať SKUEV####.sqlite a priečinok
foto\\SKUEV####.

Dátumy odovzdania (stĺpce mapovatelOK, daphneOK, sopOK zo skuev_mapovatel.txt)
sa pri každom spustení (aj --list) zapíšu do data\\datumy.json, z ktorého ich
číta zoznam území. Po zmene dátumov stačí spustiť --list a nahrať tento súbor.
Územie bez dátumu mapovatelOK sa v zozname území (index.php) nezobrazí.

Čo sa ukladá (každá databáza územia má tieto tabuľky):
  skuev      – jeden riadok o území (názov z N2000, počet polygónov, plocha,
               rozsah mapy, dátumy, mapovatelia)
  polygon    – riadky tblHabHlavna s geometriou prepočítanou z EPSG:5514 do
               WGS84 (JSON zoznam prstencov, priamo použiteľný pre Google Maps)
  biotop     – tblHabBiotopy s rozpísanými kódmi (2002 aj 2023) a kvalitou
  opatrenie  – tblHabBiotopyOpatrenia (kód opatrenia rozpísaný)
  druh       – tblHabDruhy (názov taxónu z tblHabDruhyLookup)
  aktivita   – tblAktivity (kód rozpísaný z tblAktivityLookup / AktivityLookup)
  foto       – fotky z priečinka "f" (súbor sa skopíruje do data\\foto)
  predmet_ochrany – biotopy, ktoré sú predmetom ochrany územia
               (z predmety_ochrany_SKUEV.xlsx, hárok "biotopy")
  natura_prevod   – prevod kódov katalógu 2023 / 2002 na kód Natura 2000
               (z botany_app.sqlite aplikácie bot)
  ciselnik   – kódy biotopov (PrevodKatalogy.gpkg), taxóny, aktivity
               a opatrenia, z ktorých web ponúka hodnoty pri úprave atribútov
  Riadky biotop, opatrenie, druh a aktivita majú [gpkg_fid] = [fid] v GPKG,
  aby sa úprava z webu dala zapísať späť (09_import_zmeny.py).
  Web tak na serveri nepotrebuje ani xlsx, ani katalóg. Bez zdroja sa tabuľka
  nevytvorí a web namiesto predmetov ochrany vypíše upozornenie. Po zmene xlsx
  alebo katalógu stačí bez prestavania území spustiť:

    python 08_web_hotovo.py --predmety              # všetky územia na webe
    python 08_web_hotovo.py --predmety SKUEV0870

Úpravy atribútov z webu sa ukladajú do data\\zmeny.sqlite – tento súbor patrí
webu (na serveri sa nikdy neprepisuje ani nemaže) a skript ho nečíta. Web po
pregenerovaní databázy územia neprenesené zmeny sám znova použije (pri prvom
otvorení územia), takže sa nestratia. Do GPKG ich zapíše 09_import_zmeny.py.

Prepojené formuláre ([polygon_id_form]): polygón sa uloží so svojimi vlastnými
údajmi a s odkazom na zdrojový polygón (rovnaké pravidlo ako v 06 – hľadá sa
podľa [RECORDID], záložne podľa [polygon_id], reťazenie sa sleduje až po
koncový zdroj s ochranou proti cyklu). Web pri takom polygóne zobrazí údaje
zdrojového formulára a označí ich.

Fotky: polygón sa berie z najnovšej tabuľky SKUEV####_FotoTable_*.txt
(04_validacia_foto.py) – rovnako ako v 06_naplnenie_access.py. Obrázok, ktorý
v tabuľke nie je (pribudol neskôr), sa skúsi priradiť podľa prvého čísla
v názve = [RECORDID]; nepriradená fotka sa uloží s prázdnym [polygon_fk].
Fotka patrí polygónu, pri ktorom bola odfotená – prepojený polygón svoje fotky
väčšinou nemá, web ich zobrazí cez [zdroj_polygon_fk].

Vyžaduje: shapely, pyproj (Python314). Pillow je nepovinné – bez neho sa fotky
kopírujú v pôvodnej veľkosti a bez náhľadov.
"""

import argparse
import csv
import glob
import json
import os
import re
import shutil
import sqlite3
import sys
from datetime import datetime

import shapely.wkb
from pyproj import Transformer
from shapely.ops import transform as shapely_transform

try:
    from PIL import Image, ImageOps
except ImportError:   # bez Pillow sa fotky len skopírujú
    Image = ImageOps = None

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

CLOUD_DIR = r"C:\Users\RASLAS\QField\cloud"
GPKG_NAME = "MapovaniePrePS.gpkg"
AKTIVITY_GPKG = "AktivityLookup.gpkg"
PREVOD_GPKG = "PrevodKatalogy.gpkg"   # kódy biotopov 2002/2023 (ako vo formulári)
WEB_DIR = r"C:\wamp64\www\MapovaniePrePSHotovo"
WEB_DATA_DIR = os.path.join(WEB_DIR, "data")
OLD_WEB_DB = "hotovo.sqlite"     # pôvodná spoločná databáza – už sa nepoužíva
N2000_GPKG = r"C:\_projects\MapovaniePrePS\N2000_2024.gpkg"
N2000_TABLE = "natura2000_end2024_mapovaniePrePS"
# predmety ochrany všetkých ÚEV (hárok "biotopy": SKUEV | kod | predmet_ochrany)
PREDMETY_XLSX = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                             "predmety_ochrany_SKUEV.xlsx")
PREDMETY_SHEET = "biotopy"
# katalógy biotopov s ekvivalentmi Natura 2000 (aplikácia bot)
KATALOG_DB = r"C:\wamp64\www\bot\data\botany_app.sqlite"
MAPOVATEL_FILE = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                              "skuev_mapovatel.txt")
DATUMY_COLUMNS = ("mapovatelOK", "daphneOK", "sopOK")
DATUMY_JSON = "datumy.json"      # dátumy odovzdania pre index.php
SHOW_COLUMN = "mapovatelOK"      # bez tohto dátumu index.php územie nezobrazí

# fotky: <cloud>\SKUEV####\f  ->  <priečinok data>\foto\SKUEV####
FOTO_SUBDIR = "f"                # priečinok s fotkami v cloude (ako v 07)
WEB_FOTO_DIR = "foto"            # priečinok vedľa databáz území
NAHLAD_SUBDIR = "n"              # náhľady v priečinku územia
IMAGE_EXTS = {".jpg", ".jpeg", ".png", ".heic", ".heif", ".tif", ".tiff", ".webp"}
FOTO_MAX_PX = 1600               # dlhšia strana veľkej fotky na webe
NAHLAD_MAX_PX = 320              # dlhšia strana náhľadu
FOTO_QUALITY = 85

SOURCE_CRS = 5514
TARGET_CRS = 4326
COORD_PRECISION = 6  # ~0,1 m, viac nemá zmysel

# kategórie s percentom -> text do InfoWindow
KVALITA = [("kvalita_biotopu_good", "Dobrá"),
           ("kvalita_biotopu_bad", "Zlá"),
           ("kvalita_biotopu_unsiut", "Nevyhovujúca")]
MANAZMENT = [("manazment_biotopu_vhod", "Vhodný"),
             ("manazment_biotopu_nevhod", "Nevhodný")]
VYHLIADKY = [("vyhliadky_biotopu_good", "Dobré"),
             ("vyhliadky_biotopu_bad", "Zlé"),
             ("vyhliadky_biotopu_unsiut", "Nevyhovujúce")]

SCHEMA = """
CREATE TABLE IF NOT EXISTS skuev (
    kod              TEXT PRIMARY KEY,
    nazov            TEXT,
    pocet_polygonov  INTEGER,
    plocha_ha        REAL,
    datum_od         TEXT,
    datum_do         TEXT,
    mapovatelia      TEXT,
    pocet_fotiek     INTEGER,
    lat_min REAL, lat_max REAL, lng_min REAL, lng_max REAL,
    zdroj            TEXT,
    aktualizovane    TEXT
);
CREATE TABLE IF NOT EXISTS polygon (
    id                  INTEGER PRIMARY KEY,
    skuev               TEXT NOT NULL,
    fid                 INTEGER,
    recordid            INTEGER,
    polygon_id          INTEGER,   -- číslo polygónu na webe = [RECORDID]
    typ_polygon         TEXT,
    podlaorta           TEXT,
    lokalita            TEXT,
    datum               TEXT,
    hlavny_mapovatel    TEXT,
    druhy_mapovatel     TEXT,
    p                   INTEGER,
    plocha_m2           REAL,
    poznamka            TEXT,
    e0 INTEGER, e1 INTEGER, e2 INTEGER, e3 INTEGER,
    e1_invaz INTEGER, e2_invaz INTEGER, e3_invaz INTEGER,
    polygon_id_form     INTEGER,
    zdroj_polygon_fk    INTEGER,
    zdroj_polygon_id    INTEGER,   -- [RECORDID] zdrojového polygónu
    biotopy_suhrn       TEXT,
    lat REAL, lng REAL,
    lat_min REAL, lat_max REAL, lng_min REAL, lng_max REAL,
    prstence            TEXT
);
CREATE TABLE IF NOT EXISTS biotop (
    id          INTEGER PRIMARY KEY,
    polygon_fk  INTEGER NOT NULL,
    poradie     INTEGER,
    kod2002     TEXT,
    nazov2002   TEXT,
    kod2023     TEXT,
    nazov2023   TEXT,
    pokryv      INTEGER,
    kvalita     TEXT,
    manazment   TEXT,
    vyhliadky   TEXT,
    gpkg_fid    INTEGER,   -- [fid] v tblHabBiotopy (úpravy z webu)
    -- pôvodné percentá (zdroj textov kvalita/manazment/vyhliadky – úpravy z webu)
    kvalita_biotopu_good INTEGER, kvalita_biotopu_bad INTEGER,
    kvalita_biotopu_unsiut INTEGER,
    manazment_biotopu_vhod INTEGER, manazment_biotopu_nevhod INTEGER,
    vyhliadky_biotopu_good INTEGER, vyhliadky_biotopu_bad INTEGER,
    vyhliadky_biotopu_unsiut INTEGER
);
CREATE TABLE IF NOT EXISTS opatrenie (
    id          INTEGER PRIMARY KEY,
    biotop_fk   INTEGER NOT NULL,
    kod         TEXT,
    nazov       TEXT,
    opis        TEXT,
    percento    INTEGER,
    gpkg_fid    INTEGER    -- [fid] v tblHabBiotopyOpatrenia
);
CREATE TABLE IF NOT EXISTS druh (
    id                INTEGER PRIMARY KEY,
    polygon_fk        INTEGER NOT NULL,
    kod               INTEGER,
    nazov_lat         TEXT,
    taxon             TEXT,
    pokryvnost        TEXT,
    etaz              TEXT,
    pokryvnost_perc   INTEGER,
    kod_kb            TEXT,
    charakteristicky  TEXT,
    gpkg_fid          INTEGER    -- [fid] v tblHabDruhy
);
CREATE TABLE IF NOT EXISTS aktivita (
    id           INTEGER PRIMARY KEY,
    polygon_fk   INTEGER NOT NULL,
    kod          TEXT,
    nazov        TEXT,
    intenzita    TEXT,
    perc_plochy  INTEGER,
    vplyv        TEXT,
    gpkg_fid     INTEGER    -- [fid] v tblAktivity
);
CREATE TABLE IF NOT EXISTS foto (
    id           INTEGER PRIMARY KEY,
    skuev        TEXT NOT NULL,
    polygon_fk   INTEGER,          -- prázdne = fotku sa nepodarilo priradiť
    poradie      INTEGER,
    subor        TEXT,             -- názov v data/foto/<skuev>/
    nahlad       TEXT,             -- názov v data/foto/<skuev>/n/ (môže chýbať)
    original     TEXT,             -- pôvodný názov v cloude
    poznamka     TEXT,             -- poznámka z FotoTable (ako sa priradila)
    bajtov       INTEGER
);
-- číselníky pre úpravy z webu (kódy biotopov, taxóny, aktivity, opatrenia)
CREATE TABLE IF NOT EXISTS ciselnik (
    typ    TEXT NOT NULL,   -- biotop2023, biotop2002, taxon, aktivita, opatrenie
    kod    TEXT NOT NULL,
    nazov  TEXT,
    PRIMARY KEY (typ, kod)
);
CREATE INDEX IF NOT EXISTS ix_polygon_skuev ON polygon(skuev);
CREATE INDEX IF NOT EXISTS ix_polygon_fid ON polygon(fid);
CREATE INDEX IF NOT EXISTS ix_biotop_polygon ON biotop(polygon_fk);
CREATE INDEX IF NOT EXISTS ix_opatrenie_biotop ON opatrenie(biotop_fk);
CREATE INDEX IF NOT EXISTS ix_druh_polygon ON druh(polygon_fk);
CREATE INDEX IF NOT EXISTS ix_aktivita_polygon ON aktivita(polygon_fk);
CREATE INDEX IF NOT EXISTS ix_foto_polygon ON foto(polygon_fk);
CREATE INDEX IF NOT EXISTS ix_foto_skuev ON foto(skuev);
CREATE INDEX IF NOT EXISTS ix_biotop_gpkg ON biotop(gpkg_fid);
CREATE INDEX IF NOT EXISTS ix_opatrenie_gpkg ON opatrenie(gpkg_fid);
CREATE INDEX IF NOT EXISTS ix_druh_gpkg ON druh(gpkg_fid);
CREATE INDEX IF NOT EXISTS ix_aktivita_gpkg ON aktivita(gpkg_fid);
"""

# vytvoria sa len vtedy, keď sa zdroj podarilo prečítať (inak web vypíše chybu)
PREDMETY_SCHEMA = """
CREATE TABLE predmet_ochrany (
    kod    TEXT PRIMARY KEY,   -- kód Natura 2000, napr. 6510, 91E0
    nazov  TEXT
);
"""
NATURA_SCHEMA = """
CREATE TABLE natura_prevod (
    katalog  TEXT NOT NULL,    -- 'katalog' (2023) alebo 'katalog2002'
    kod      TEXT NOT NULL,    -- kód biotopu v katalógu
    natura   TEXT NOT NULL,    -- kód Natura 2000
    PRIMARY KEY (katalog, kod)
);
"""


# ----------------------------------------------------------------------
# Pomocné funkcie
# ----------------------------------------------------------------------

def normalize(code):
    """'870' / '0870' / 'skuev0870' -> 'SKUEV0870', inak None."""
    code = str(code).strip().upper()
    if re.fullmatch(r"\d{1,4}", code):
        code = "SKUEV" + code.zfill(4)
    return code if re.fullmatch(r"SKUEV\d{4}", code) else None


def find_folder(skuev):
    """Len veľkopísmenový priečinok SKUEV#### (rovnako ako 06)."""
    try:
        dir_names = os.listdir(CLOUD_DIR)
    except OSError as e:
        sys.exit("Priečinok %s sa nedá čítať: %s" % (CLOUD_DIR, e))
    if skuev not in dir_names:
        sys.exit("Priečinok %s neexistuje v %s (malé 'skuev...' sa ignoruje)."
                 % (skuev, CLOUD_DIR))
    return os.path.join(CLOUD_DIR, skuev)


def gpkg_wkb(blob):
    """WKB časť z GPKG geometrie, alebo None."""
    if blob is None or len(blob) < 8 or blob[0:2] != b"GP":
        return None
    env_counts = {0: 0, 1: 4, 2: 6, 3: 6, 4: 8}
    env_code = (blob[3] >> 1) & 0x07
    if env_code not in env_counts:
        return None
    return blob[8 + 8 * env_counts[env_code]:]


def geom_from_blob(blob):
    wkb = gpkg_wkb(blob)
    if not wkb:
        return None
    try:
        return shapely.wkb.loads(wkb)
    except Exception:
        return None


def is_empty(v):
    return v is None or (isinstance(v, str) and not v.strip())


def clean(v):
    """Text bez okrajových medzier; prázdny -> None."""
    if v is None:
        return None
    if isinstance(v, str):
        v = v.strip()
        return v or None
    return v


def rings_wgs84(geom, transformer):
    """Zoznam prstencov [[lng, lat], ...] v WGS84 (aj diery, aj multipolygóny)."""
    geom = shapely_transform(
        lambda x, y, z=None: transformer.transform(x, y), geom)
    rings = []
    parts = list(geom.geoms) if geom.geom_type.startswith("Multi") else [geom]
    for part in parts:
        if part.is_empty:
            continue
        for ring in [part.exterior] + list(part.interiors):
            rings.append([[round(x, COORD_PRECISION), round(y, COORD_PRECISION)]
                          for x, y in ring.coords])
    return rings


def percenta(row, categories):
    """'Dobrá 50 %, Nevyhovujúca 50 %' z percentuálnych stĺpcov, inak None."""
    parts = []
    for column, label in categories:
        value = row[column]
        if value:
            parts.append("%s %g %%" % (label, value))
    return ", ".join(parts) or None


def load_lookup(con, table, key_col, value_col):
    """{kľúč: hodnota} z lookup tabuľky; chýbajúca tabuľka -> prázdny slovník."""
    try:
        rows = con.execute('SELECT "%s", "%s" FROM "%s"'
                           % (key_col, value_col, table)).fetchall()
    except sqlite3.Error:
        return {}
    return {str(k).strip(): v for k, v in rows if k is not None}


def site_name(skuev):
    """SITENAME z N2000; ak sa nedá zistiť, None."""
    if not os.path.isfile(N2000_GPKG):
        return None
    try:
        con = sqlite3.connect(N2000_GPKG)
        row = con.execute('SELECT SITENAME FROM "%s" WHERE SITECODE = ?'
                          % N2000_TABLE, (skuev,)).fetchone()
        con.close()
        return row[0] if row else None
    except sqlite3.Error:
        return None


def cell_text(value):
    """Hodnota bunky xlsx ako text (kód 6510 môže byť uložený ako číslo)."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def read_predmety():
    """Predmety ochrany všetkých území: {SKUEV####: {kod: názov}}.

    Ak sa xlsx nedá prečítať, vráti (None, popis chyby)."""
    try:
        import openpyxl
    except ImportError:
        return None, "chýba balík openpyxl (pip install openpyxl)"
    try:
        wb = openpyxl.load_workbook(PREDMETY_XLSX, read_only=True, data_only=True)
    except Exception as e:
        return None, "%s sa nedá otvoriť (%s)" % (PREDMETY_XLSX, e)
    try:
        if PREDMETY_SHEET not in wb.sheetnames:
            return None, "v %s chýba hárok \"%s\"" % (PREDMETY_XLSX, PREDMETY_SHEET)
        predmety = {}
        for row in wb[PREDMETY_SHEET].iter_rows(min_row=2, values_only=True):
            row = tuple(row) + (None, None, None)
            skuev, kod, nazov = (cell_text(v) for v in row[:3])
            if normalize(skuev) and kod:
                predmety.setdefault(normalize(skuev), {})[kod] = nazov
        return predmety, None
    finally:
        wb.close()


def read_natura_prevod():
    """Prevod kódov katalógov na kód Natura 2000: [(katalog, kod, natura)].

    Nadradený biotop bez vlastného kódu preberie kód podjednotiek, ak majú
    všetky rovnaký. Ak sa katalóg nedá prečítať, vráti (None, popis chyby)."""
    if not os.path.isfile(KATALOG_DB):
        return None, "katalóg biotopov %s neexistuje" % KATALOG_DB
    try:
        con = sqlite3.connect("file:%s?mode=ro" % KATALOG_DB.replace("\\", "/"),
                              uri=True)
        rows = con.execute(
            "SELECT source, code, details_json FROM biotopes "
            "WHERE source IN ('katalog', 'katalog2002')").fetchall()
        con.close()
    except sqlite3.Error as e:
        return None, "katalóg biotopov %s sa nedá prečítať (%s)" % (KATALOG_DB, e)
    kody = {"katalog": {}, "katalog2002": {}}
    deti = {}
    for source, code, details in rows:
        try:
            d = json.loads(details or "{}")
        except ValueError:
            d = {}
        for c in d.get("classifications") or []:
            if c.get("system") == "Natura 2000" and c.get("code"):
                kody[source][code] = c["code"]
                break
        deti.setdefault(source, {})[code] = [
            ch.get("code") for ch in d.get("children") or []]
    for source, by_code in deti.items():
        for code, children in by_code.items():
            if code in kody[source] or not children:
                continue
            z_deti = {kody[source].get(ch, "") for ch in children}
            if len(z_deti) == 1 and "" not in z_deti:
                kody[source][code] = z_deti.pop()
    return [(source, code, natura) for source, by_code in kody.items()
            for code, natura in by_code.items()], None


def write_predmety(con, skuev, predmety, prevod):
    """Zapíše tabuľky predmet_ochrany a natura_prevod (nanovo).

    `predmety` / `prevod` sú výsledky read_predmety / read_natura_prevod;
    tabuľka, ktorej zdroj chýba (None), sa nevytvorí. Vráti počet predmetov
    ochrany územia, alebo None, ak sa nezapísali. Commit robí volajúci."""
    # jedna transakcia – web počas prepisovania nevidí chýbajúcu tabuľku
    if not con.in_transaction:
        con.execute("BEGIN")
    con.execute("DROP TABLE IF EXISTS predmet_ochrany")
    con.execute("DROP TABLE IF EXISTS natura_prevod")
    if prevod is not None:
        con.execute(NATURA_SCHEMA)
        con.executemany("INSERT INTO natura_prevod (katalog, kod, natura) "
                        "VALUES (?,?,?)", prevod)
    if predmety is None:
        return None
    con.execute(PREDMETY_SCHEMA)
    site_predmety = predmety.get(skuev, {})
    con.executemany("INSERT INTO predmet_ochrany (kod, nazov) VALUES (?,?)",
                    sorted(site_predmety.items()))
    return len(site_predmety)


def parse_datum(value):
    """'dd.mm.rrrr' -> date, inak None (na zoradenie dátumov mapovania)."""
    if is_empty(value):
        return None
    for fmt in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(str(value).strip(), fmt).date()
        except ValueError:
            continue
    return None


# ----------------------------------------------------------------------
# Fotky
# ----------------------------------------------------------------------

def foto_table(folder, skuev, warnings):
    """Údaje o fotkách z najnovšej SKUEV####_FotoTable_*.txt.

    Súbor vytvára 04_validacia_foto.py; stĺpce sa čítajú podľa hlavičky, aby
    sedeli aj staršie tabuľky (v prvom stĺpci majú [polygon_id] namiesto
    [RECORDID]) – rovnako ako v 06_naplnenie_access.py:
        RECORDID | polygon_id, original_name, [new_name], note
    Vráti ({názov súboru malými písmenami: (kľúč | None, poznámka | None)},
    kľúč_je_recordid, cesta_k_tabuľke | None).
    """
    files = glob.glob(os.path.join(folder, skuev + "_FotoTable_*.txt"))
    if not files:
        return {}, True, None
    latest = max(files, key=os.path.getmtime)
    info = {}
    try:
        with open(latest, encoding="utf-8") as f:
            header = [h.strip().lower() for h in f.readline().rstrip("\n").split("\t")]
            col = {name: i for i, name in enumerate(header)}
            by_recordid = "recordid" in col
            key_idx = col.get("recordid", col.get("polygon_id", 0))
            name_idx = col.get("original_name")
            note_idx = col.get("note")
            if name_idx is None:
                warnings.append("%s nemá stĺpec 'original_name' – tabuľka sa "
                                "ignoruje." % os.path.basename(latest))
                return {}, by_recordid, latest
            def value(parts, i):
                return "" if i is None or i >= len(parts) else parts[i].strip()

            for line in f:
                parts = line.rstrip("\n").split("\t")
                name = value(parts, name_idx)
                if not name:
                    continue  # riadok "polygón bez fotky" – žiadny súbor
                info[name.lower()] = (value(parts, key_idx) or None,
                                      value(parts, note_idx) or None)
    except OSError as e:
        warnings.append("%s sa nedá prečítať (%s) – fotky sa priradia podľa "
                        "názvov súborov." % (latest, e))
        return {}, True, latest
    return info, by_recordid, latest


def read_fotky(folder, skuev, rec_map, pid_map, warnings):
    """Fotky územia z priečinka "f".

    `rec_map` je {str([RECORDID]): [RECORDID]} existujúcich polygónov,
    `pid_map` je {str([polygon_id]): [RECORDID]} pre staršie FotoTable.
    Vráti ({RECORDID: [fotky]}, [nepriradené fotky]); fotka je slovník
    s kľúčmi 'cesta', 'original' a 'poznamka'.
    """
    foto_dir = os.path.join(folder, FOTO_SUBDIR)
    if not os.path.isdir(foto_dir):
        warnings.append("Priečinok s fotkami %s neexistuje – územie ostane bez "
                        "fotiek." % foto_dir)
        return {}, []
    files = sorted(n for n in os.listdir(foto_dir)
                   if os.path.splitext(n)[1].lower() in IMAGE_EXTS
                   and os.path.isfile(os.path.join(foto_dir, n)))
    if not files:
        warnings.append("V %s nie sú žiadne obrázky." % foto_dir)
        return {}, []

    info, by_recordid, table_path = foto_table(folder, skuev, warnings)
    if table_path is None:
        warnings.append("Nenašiel som %s_FotoTable_*.txt – polygón fotky sa určí "
                        "len podľa čísla v názve súboru (spusti "
                        "04_validacia_foto.py)." % skuev)
    key_map = rec_map if by_recordid else pid_map

    by_rec, nepriradene = {}, []
    for name in files:
        if name.lower() in info:
            key, note = info[name.lower()]
            rec = key_map.get(key) if key else None
        else:
            # súbor pribudol až po behu 04 – prvé číslo v názve = [RECORDID]
            m = re.search(r"\d+", os.path.splitext(name)[0])
            key = m.group(0) if m else None
            rec = rec_map.get(key) if key else None
            note = "nie je vo FotoTable – polygón podľa názvu súboru"
            warnings.append("%s nie je vo FotoTable (pribudla neskôr?) – %s."
                            % (name, "priradená podľa názvu" if rec is not None
                               else "ostáva nepriradená"))
        foto = {"cesta": os.path.join(foto_dir, name),
                "original": name,
                "poznamka": note}
        if rec is None:
            if key:
                warnings.append("%s: kľúč '%s' nezodpovedá žiadnemu polygónu "
                                "územia – fotka ostane nepriradená." % (name, key))
            nepriradene.append(foto)
        else:
            by_rec.setdefault(rec, []).append(foto)
    return by_rec, nepriradene


def safe_base(name):
    """Názov súboru bez prípony, bezpečný pre URL ('IMG 1.JPG' -> 'IMG_1')."""
    base = os.path.splitext(os.path.basename(name))[0]
    base = re.sub(r"[^A-Za-z0-9._-]+", "_", base).strip("._-")
    return base or "foto"


def web_copy(src, dst_dir, base, max_px, warnings):
    """Zmenšená JPEG kópia obrázka v `dst_dir`; vráti názov výsledného súboru.

    Bez Pillow (alebo keď zmenšenie zlyhá) sa súbor len skopíruje. EXIF sa
    zámerne nezachováva – obrázok sa otočí podľa orientácie a na web sa tak
    nedostanú ani GPS súradnice z fotoaparátu.
    """
    if Image is not None:
        dst = os.path.join(dst_dir, base + ".jpg")
        try:
            with Image.open(src) as img:
                img = ImageOps.exif_transpose(img)
                img.thumbnail((max_px, max_px), Image.LANCZOS)
                if img.mode not in ("RGB", "L"):
                    img = img.convert("RGB")
                img.save(dst, "JPEG", quality=FOTO_QUALITY, optimize=True)
            return base + ".jpg"
        except Exception as e:
            warnings.append("%s: zmenšenie zlyhalo (%s) – kopírujem originál."
                            % (os.path.basename(src), e))
            if os.path.exists(dst):
                os.remove(dst)
    name = base + os.path.splitext(src)[1].lower()
    shutil.copy2(src, os.path.join(dst_dir, name))
    return name


def remove_foto_dir(foto_root, skuev, warnings=None):
    """Zmaže priečinok s fotkami jedného územia (ak existuje)."""
    if not re.fullmatch(r"SKUEV\d{4}", skuev):   # poistka pred rmtree
        return
    site_dir = os.path.join(foto_root, skuev)
    if not os.path.isdir(site_dir):
        return
    try:
        shutil.rmtree(site_dir)
    except OSError as e:
        message = "Staré fotky v %s sa nepodarilo zmazať (%s)." % (site_dir, e)
        if warnings is None:
            print("   UPOZORNENIE: %s" % message)
        else:
            warnings.append(message)


def export_fotky(site, polygons, foto_root, warnings):
    """Skopíruje fotky územia do <foto_root>\\SKUEV#### (+ náhľady do "n").

    Každej fotke doplní 'subor', 'nahlad' a 'bajtov' – to, čo sa zapíše do
    databázy. Staré fotky územia sa vždy prepíšu. Vráti (počet, veľkosť v B).
    """
    remove_foto_dir(foto_root, site["kod"], warnings)
    vsetky = ([f for p in polygons for f in p["fotky"]] + site["fotky_nepriradene"])
    if not vsetky:
        return 0, 0
    if Image is None:
        warnings.append("Pillow nie je nainštalované – fotky sa kopírujú "
                        "v pôvodnej veľkosti a bez náhľadov.")

    site_dir = os.path.join(foto_root, site["kod"])
    nahlad_dir = os.path.join(site_dir, NAHLAD_SUBDIR)
    os.makedirs(nahlad_dir if Image is not None else site_dir, exist_ok=True)

    pouzite = set()
    pocet = bajtov = 0
    for foto in vsetky:
        base = safe_base(foto["original"])
        unikat, poradie = base, 1
        while unikat.lower() in pouzite:
            poradie += 1
            unikat = "%s_%d" % (base, poradie)
        pouzite.add(unikat.lower())
        try:
            foto["subor"] = web_copy(foto["cesta"], site_dir, unikat,
                                     FOTO_MAX_PX, warnings)
            foto["nahlad"] = (web_copy(foto["cesta"], nahlad_dir, unikat,
                                       NAHLAD_MAX_PX, warnings)
                              if Image is not None else None)
        except OSError as e:
            warnings.append("%s: fotku sa nepodarilo skopírovať (%s)."
                            % (foto["original"], e))
            continue
        foto["bajtov"] = os.path.getsize(os.path.join(site_dir, foto["subor"]))
        bajtov += foto["bajtov"]
        pocet += 1
    return pocet, bajtov


# ----------------------------------------------------------------------
# Načítanie jedného územia z GPKG
# ----------------------------------------------------------------------

def read_ciselnik(folder, biotopy2002, biotopy2023, taxony,
                  opatrenia_lookup, aktivity_lookup, warnings):
    """Číselníky pre úpravy z webu: [(typ, kod, nazov)].

    Kódy biotopov sú tie, ktoré ponúka formulár (PrevodKatalogy, first_20xx = 1);
    názov sa berie z lookupu GPKG ako pri zobrazení, inak z prevodníka."""
    rows = {}

    def add(typ, kod, nazov):
        kod = clean(kod)
        if kod is None or kod == "-":
            return
        rows.setdefault((typ, str(kod)), clean(nazov))

    prevod = os.path.join(folder, PREVOD_GPKG)
    if not os.path.isfile(prevod):
        prevod = os.path.join(os.path.dirname(os.path.abspath(__file__)), PREVOD_GPKG)
    try:
        pcon = sqlite3.connect("file:%s?mode=ro" % prevod.replace("\\", "/"), uri=True)
        for kod, nazov in pcon.execute(
                "SELECT kod_2023, nazov_2023 FROM PrevodKatalogy WHERE first_2023 = 1"):
            add("biotop2023", kod, biotopy2023.get(clean(kod) or "") or nazov)
        for kod, nazov in pcon.execute(
                "SELECT kod_2002, nazov_2002 FROM PrevodKatalogy WHERE first_2002 = 1"):
            add("biotop2002", kod, biotopy2002.get(clean(kod) or "") or nazov)
        pcon.close()
    except sqlite3.Error as e:
        warnings.append("%s sa nedá čítať (%s) – kódy biotopov pre úpravy z webu "
                        "sa vezmú z lookupov GPKG." % (prevod, e))
        for kod, nazov in biotopy2023.items():
            add("biotop2023", kod, nazov)
        for kod, nazov in biotopy2002.items():
            add("biotop2002", kod, nazov)
    for kod, nazov in taxony.items():
        add("taxon", kod, nazov)
    for kod, nazov in opatrenia_lookup.items():
        add("opatrenie", kod, nazov)
    for kod, nazov in aktivity_lookup.items():
        add("aktivita", kod, nazov)
    return [(typ, kod, nazov) for (typ, kod), nazov in sorted(rows.items())]


def read_site(skuev, folder, warnings, s_fotkami=True):
    """Prečíta údaje územia z GPKG a vráti (skuev_riadok, polygóny).

    Nepriradené fotky sú v skuev_riadku pod kľúčom 'fotky_nepriradene'."""
    gpkg = os.path.join(folder, GPKG_NAME)
    if not os.path.isfile(gpkg):
        sys.exit("Súbor %s neexistuje." % gpkg)

    con = sqlite3.connect(gpkg)
    con.row_factory = sqlite3.Row

    # lookupy: biotopy, taxóny, opatrenia (v GPKG) + aktivity (vedľajší GPKG)
    biotopy2002 = load_lookup(con, "tblHabBiotopyLookup", "code", "biotop_name")
    biotopy2023 = load_lookup(con, "tblHabBiotopyNewLookup", "codenew",
                              "biotopnew_name")
    taxony = load_lookup(con, "tblHabDruhyLookup", "Tax_id", "Taxon_meno")
    opatrenia_lookup = load_lookup(con, "tblAktivityLookup", "node_code", "namex")
    aktivity_lookup = dict(opatrenia_lookup)
    akt_gpkg = os.path.join(folder, AKTIVITY_GPKG)
    if os.path.isfile(akt_gpkg):
        akt_con = sqlite3.connect(akt_gpkg)
        aktivity_lookup.update(load_lookup(akt_con, "AktivityLookup", "kod", "namex"))
        akt_con.close()
    else:
        warnings.append("%s chýba – kódy aktivít ostanú nerozpísané." % AKTIVITY_GPKG)

    hlavna = con.execute("SELECT * FROM tblHabHlavna").fetchall()
    cudzie = [r for r in hlavna
              if not is_empty(r["KOD_UEV"]) and str(r["KOD_UEV"]).strip() != skuev]
    if cudzie:
        warnings.append("%d polygónov má iný [KOD_UEV] než %s – vynechané (%s)."
                        % (len(cudzie), skuev,
                           ", ".join(sorted({str(r["KOD_UEV"]) for r in cudzie}))))
    # polygóny dokreslené v teréne nemajú [KOD_UEV] – gpkg patrí jednému územiu,
    # takže k nemu patria aj ony
    hlavna = [r for r in hlavna
              if is_empty(r["KOD_UEV"]) or str(r["KOD_UEV"]).strip() == skuev]
    if not hlavna:
        sys.exit("V %s nie je ani jeden polygón s [KOD_UEV] = '%s'." % (gpkg, skuev))

    # --- prepojené formuláre: rovnaké pravidlo ako v 06_naplnenie_access.py ---
    by_rec, by_pid = {}, {}
    for r in sorted(hlavna, key=lambda x: x["fid"]):
        if not is_empty(r["RECORDID"]):
            by_rec.setdefault(str(r["RECORDID"]).strip(), r)
        if not is_empty(r["polygon_id"]):
            by_pid.setdefault(str(r["polygon_id"]).strip(), r)

    def polyname(row):
        return "RECORDID=%s (fid=%s)" % (row["RECORDID"], row["fid"])

    # --- fotky: kľúč FotoTable ([RECORDID], staršie [polygon_id]) -> RECORDID ---
    if s_fotkami:
        rec_map = {rec: rec for rec in
                   (str(r["RECORDID"]).strip() for r in hlavna
                    if not is_empty(r["RECORDID"]))}
        pid_map = {pid: str(row["RECORDID"]).strip()
                   for pid, row in by_pid.items() if not is_empty(row["RECORDID"])}
        fotky_by_rec, fotky_nepriradene = read_fotky(folder, skuev, rec_map,
                                                     pid_map, warnings)
    else:
        fotky_by_rec, fotky_nepriradene = {}, []

    def source_row(row):
        current = row
        seen = {row["fid"]}
        while not is_empty(current["polygon_id_form"]):
            ref = str(current["polygon_id_form"]).strip()
            parent = by_rec.get(ref) or by_pid.get(ref)
            if parent is None:
                warnings.append("%s: [polygon_id_form]='%s' nemá zodpovedajúci "
                                "polygón – použité vlastné údaje."
                                % (polyname(row), ref))
                return current
            if parent["fid"] in seen:
                warnings.append("%s: cyklus v [polygon_id_form]." % polyname(row))
                return current
            seen.add(parent["fid"])
            current = parent
        return current

    # --- deti podľa fkRECORDID ---
    def index_by_rec(table):
        idx = {}
        try:
            rows = con.execute("SELECT * FROM %s" % table).fetchall()
        except sqlite3.Error:
            return idx
        for row in rows:
            if not is_empty(row["fkRECORDID"]):
                idx.setdefault(str(row["fkRECORDID"]).strip(), []).append(row)
        return idx

    druhy_by_rec = index_by_rec("tblHabDruhy")
    biotopy_by_rec = index_by_rec("tblHabBiotopy")
    aktivity_by_rec = index_by_rec("tblAktivity")

    # opatrenia sa viažu na tblHabBiotopy.[id], ktoré nie je jedinečné –
    # rovnako ako v 06 vyhráva biotop s najnižším fid
    biotop_fid_by_id = {}
    ambiguous = set()
    for b in sorted((r for rows in biotopy_by_rec.values() for r in rows),
                    key=lambda x: x["fid"]):
        if b["id"] in biotop_fid_by_id:
            ambiguous.add(b["id"])
        else:
            biotop_fid_by_id[b["id"]] = b["fid"]
    opatrenia_by_biotop = {}
    try:
        opat_rows = con.execute("SELECT * FROM tblHabBiotopyOpatrenia").fetchall()
    except sqlite3.Error:
        opat_rows = []
    for o in opat_rows:
        bfid = biotop_fid_by_id.get(o["fkHabBiotopyID"])
        if bfid is None:
            warnings.append("tblHabBiotopyOpatrenia fid=%s: [fkHabBiotopyID]=%s "
                            "nemá biotop – vynechané."
                            % (o["fid"], o["fkHabBiotopyID"]))
            continue
        if o["fkHabBiotopyID"] in ambiguous:
            warnings.append("tblHabBiotopyOpatrenia fid=%s: [fkHabBiotopyID]=%s "
                            "zodpovedá viacerým biotopom – priradené biotopu fid=%s."
                            % (o["fid"], o["fkHabBiotopyID"], bfid))
        opatrenia_by_biotop.setdefault(bfid, []).append(o)

    transformer = Transformer.from_crs(SOURCE_CRS, TARGET_CRS, always_xy=True)
    polygons = []
    plocha_celkom = 0.0
    datumy = []
    mapovatelia = []
    bbox = [None, None, None, None]  # lat_min, lat_max, lng_min, lng_max

    def recordid_key(row):
        try:
            return int(row["RECORDID"])
        except (TypeError, ValueError):
            return 0

    for r in sorted(hlavna, key=lambda x: (recordid_key(x), x["fid"])):
        geom = geom_from_blob(r["geom"])
        if geom is None or geom.is_empty:
            warnings.append("%s: chýba alebo je poškodená geometria – vynechané."
                            % polyname(r))
            continue
        if not geom.is_valid:
            geom = geom.buffer(0)
        rings = rings_wgs84(geom, transformer)
        if not rings:
            warnings.append("%s: geometriu sa nepodarilo previesť do WGS84."
                            % polyname(r))
            continue

        lngs = [c[0] for ring in rings for c in ring]
        lats = [c[1] for ring in rings for c in ring]
        stred = geom.representative_point()
        stred_lng, stred_lat = transformer.transform(stred.x, stred.y)

        plocha_m2 = geom.area  # EPSG:5514 je metrický
        plocha_celkom += plocha_m2
        d = parse_datum(r["datum"])
        if d:
            datumy.append(d)
        for m in (r["hlavny_mapovatel"], r["druhy_mapovatel"]):
            if not is_empty(m) and str(m).strip() not in mapovatelia:
                mapovatelia.append(str(m).strip())

        bbox = [min(x for x in (bbox[0], min(lats)) if x is not None),
                max(x for x in (bbox[1], max(lats)) if x is not None),
                min(x for x in (bbox[2], min(lngs)) if x is not None),
                max(x for x in (bbox[3], max(lngs)) if x is not None)]

        src = source_row(r)
        own_rec = None if is_empty(r["RECORDID"]) else str(r["RECORDID"]).strip()
        src_rec = None if is_empty(src["RECORDID"]) else str(src["RECORDID"]).strip()

        def z_formulara(column):
            """Hodnota formulára: zo zdrojového polygónu, prázdna -> vlastná.

            Rovnaké pravidlo ako 05_export_shapefile.py – údaje prepojeného
            formulára patria polygónu, do ktorého boli zapísané."""
            value = src[column]
            return clean(value if not is_empty(value) else r[column])

        # biotopy zoradené podľa pokryvu zostupne (ako v 05_export_shapefile)
        biotopy = sorted(biotopy_by_rec.get(src_rec) or [],
                         key=lambda b: (-(b["biotop_pokryv"] or 0), b["id"] or 0))
        biotopy_rows = []
        for poradie, b in enumerate(biotopy, start=1):
            kod2002 = clean(b["biotop_cislo"])
            kod2023 = clean(b["biotop_cislo_new"])
            biotopy_rows.append({
                "poradie": poradie,
                "kod2002": kod2002,
                "nazov2002": biotopy2002.get(kod2002) if kod2002 else None,
                "kod2023": kod2023,
                "nazov2023": biotopy2023.get(kod2023) if kod2023 else None,
                "pokryv": b["biotop_pokryv"],
                "kvalita": percenta(b, KVALITA),
                "manazment": percenta(b, MANAZMENT),
                "vyhliadky": percenta(b, VYHLIADKY),
                "gpkg_fid": b["fid"],
                "percenta": {c: b[c] for c, _ in KVALITA + MANAZMENT + VYHLIADKY},
                "opatrenia": [{
                    "kod": clean(o["kod_opatrenia"]),
                    "nazov": opatrenia_lookup.get(clean(o["kod_opatrenia"]) or ""),
                    "opis": clean(o["detailny_opis_opatrenia"]),
                    "percento": o["percento_z_plochy_biotopu"],
                    "gpkg_fid": o["fid"],
                } for o in opatrenia_by_biotop.get(b["fid"], [])],
            })
        suhrn = "; ".join(
            "%s%s" % (b["kod2023"] or b["kod2002"] or "?",
                      "" if b["pokryv"] is None else " %g %%" % b["pokryv"])
            for b in biotopy_rows) or None

        druhy_rows = [{
            "kod": d["KOD"],
            "nazov_lat": clean(d["NAZOV_LAT"]),
            "taxon": taxony.get(str(d["KOD"]).strip()) if d["KOD"] is not None else None,
            "pokryvnost": clean(d["POKRYVNOST"]),
            "etaz": clean(d["etaz"]),
            "pokryvnost_perc": d["pokryvnost_perc"],
            "kod_kb": clean(d["kod_kbx"]),
            "charakteristicky": clean(d["is_characetristic"]),
            "gpkg_fid": d["fid"],
        } for d in sorted(druhy_by_rec.get(src_rec) or [],
                          key=lambda x: (x["etaz"] or "", x["NAZOV_LAT"] or ""))]

        aktivity_rows = [{
            "kod": clean(a["Aktivita"]),
            "nazov": aktivity_lookup.get(clean(a["Aktivita"]) or ""),
            "intenzita": clean(a["Intenzita"]),
            "perc_plochy": a["Perc_Plochy"],
            "vplyv": clean(a["Vplyv"]),
            "gpkg_fid": a["fid"],
        } for a in aktivity_by_rec.get(src_rec) or []]

        polygons.append({
            "fid": r["fid"],
            "recordid": r["RECORDID"],
            # web zobrazuje [polygon_id] ako číslo polygónu – ukladá sa RECORDID
            "polygon_id": r["RECORDID"],
            # vlastné údaje polygónu (geometria, poradie, plocha, podklad)
            "podlaorta": clean(r["podlaorta"]),
            "p": r["p"],
            "plocha_m2": plocha_m2,
            # údaje formulára – pri prepojenom polygóne zo zdrojového riadku
            "typ_polygon": z_formulara("typ_polygon"),
            "lokalita": z_formulara("lokalita"),
            "datum": z_formulara("datum"),
            "hlavny_mapovatel": z_formulara("hlavny_mapovatel"),
            "druhy_mapovatel": z_formulara("druhy_mapovatel"),
            "poznamka": z_formulara("poznamka"),
            "e0": z_formulara("e0"), "e1": z_formulara("e1"),
            "e2": z_formulara("e2"), "e3": z_formulara("e3"),
            "e1_invaz": z_formulara("E1_invaz"),
            "e2_invaz": z_formulara("E2_invaz"),
            "e3_invaz": z_formulara("E3_invaz"),
            "polygon_id_form": r["polygon_id_form"],
            # vlastný riadok = None; prepojený = RECORDID zdrojového polygónu
            "zdroj_polygon_id": (src["RECORDID"] if src["fid"] != r["fid"] else None),
            "zdroj_fid": (src["fid"] if src["fid"] != r["fid"] else None),
            "biotopy_suhrn": suhrn,
            "lat": round(stred_lat, COORD_PRECISION),
            "lng": round(stred_lng, COORD_PRECISION),
            "lat_min": min(lats), "lat_max": max(lats),
            "lng_min": min(lngs), "lng_max": max(lngs),
            "prstence": json.dumps(rings, separators=(",", ":")),
            "biotopy": biotopy_rows,
            "druhy": druhy_rows,
            "aktivity": aktivity_rows,
            # fotka patrí polygónu, pri ktorom bola odfotená (vlastné RECORDID)
            "fotky": fotky_by_rec.pop(own_rec, []) if own_rec else [],
            "vlastne_data": own_rec == src_rec,
        })

    ciselnik = read_ciselnik(folder, biotopy2002, biotopy2023, taxony,
                             opatrenia_lookup, aktivity_lookup, warnings)
    con.close()

    # fotky polygónov, ktoré sa do exportu nedostali (vynechaná geometria)
    for rec, fotky in sorted(fotky_by_rec.items()):
        warnings.append("RECORDID=%s: polygón nie je v exporte, jeho %d fotiek "
                        "ostáva nepriradených." % (rec, len(fotky)))
        fotky_nepriradene.extend(fotky)

    site = {
        "kod": skuev,
        "nazov": site_name(skuev),
        "fotky_nepriradene": fotky_nepriradene,
        "pocet_polygonov": len(polygons),
        "plocha_ha": round(plocha_celkom / 10000.0, 2),
        "datum_od": min(datumy).strftime("%d.%m.%Y") if datumy else None,
        "datum_do": max(datumy).strftime("%d.%m.%Y") if datumy else None,
        "mapovatelia": ", ".join(sorted(mapovatelia)) or None,
        "lat_min": bbox[0], "lat_max": bbox[1],
        "lng_min": bbox[2], "lng_max": bbox[3],
        "zdroj": gpkg,
        "aktualizovane": datetime.now().strftime("%d.%m.%Y %H:%M"),
        "ciselnik": ciselnik,
    }
    return site, polygons


# ----------------------------------------------------------------------
# Zápis do webovej databázy
# ----------------------------------------------------------------------

def site_db_path(data_dir, skuev):
    """Databáza jedného územia: <data>\\SKUEV####.sqlite."""
    return os.path.join(data_dir, skuev + ".sqlite")


def open_web_db(path):
    con = sqlite3.connect(path)
    con.executescript(SCHEMA)
    return con


def remove_site(data_dir, skuev, foto_root=None):
    """Zmaže databázu územia (a fotky, ak je zadaný `foto_root`).

    Vráti (počet polygónov, počet fotiek), alebo None, ak územie na webe nie je."""
    path = site_db_path(data_dir, skuev)
    if foto_root:
        remove_foto_dir(foto_root, skuev)
    if not os.path.isfile(path):
        return None
    try:
        con = sqlite3.connect(path)
        counts = (con.execute("SELECT COUNT(*) FROM polygon").fetchone()[0],
                  con.execute("SELECT COUNT(*) FROM foto").fetchone()[0])
        con.close()
    except sqlite3.Error:
        counts = (0, 0)
    try:
        os.remove(path)
    except OSError as e:
        sys.exit("%s sa nedá zmazať (%s)." % (path, e))
    return counts


def insert_fotky(con, skuev, polygon_fk, fotky):
    """Vloží fotky jedného polygónu (alebo nepriradené, keď je fk None)."""
    poradie = 0
    for f in fotky:
        if not f.get("subor"):
            continue  # fotka sa nepodarila skopírovať
        poradie += 1
        con.execute(
            "INSERT INTO foto (skuev, polygon_fk, poradie, subor, nahlad, "
            "original, poznamka, bajtov) VALUES (?,?,?,?,?,?,?,?)",
            (skuev, polygon_fk, poradie, f["subor"], f.get("nahlad"),
             f["original"], f.get("poznamka"), f.get("bajtov")))
    return poradie


def write_site(con, site, polygons):
    pocet_fotiek = len([f for p in polygons for f in p["fotky"] if f.get("subor")]
                       + [f for f in site["fotky_nepriradene"] if f.get("subor")])
    con.execute(
        "INSERT INTO skuev (kod, nazov, pocet_polygonov, plocha_ha, datum_od, "
        "datum_do, mapovatelia, pocet_fotiek, lat_min, lat_max, lng_min, "
        "lng_max, zdroj, aktualizovane) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
        (site["kod"], site["nazov"], site["pocet_polygonov"], site["plocha_ha"],
         site["datum_od"], site["datum_do"], site["mapovatelia"], pocet_fotiek,
         site["lat_min"], site["lat_max"], site["lng_min"], site["lng_max"],
         site["zdroj"], site["aktualizovane"]))

    fk_by_fid = {}
    counts = {"biotop": 0, "opatrenie": 0, "druh": 0, "aktivita": 0, "foto": 0}
    for p in polygons:
        cur = con.execute(
            "INSERT INTO polygon (skuev, fid, recordid, polygon_id, typ_polygon, "
            "podlaorta, lokalita, datum, hlavny_mapovatel, druhy_mapovatel, p, "
            "plocha_m2, poznamka, e0, e1, e2, e3, e1_invaz, e2_invaz, e3_invaz, "
            "polygon_id_form, zdroj_polygon_id, biotopy_suhrn, lat, lng, "
            "lat_min, lat_max, lng_min, lng_max, prstence) "
            "VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
            (site["kod"], p["fid"], p["recordid"], p["polygon_id"], p["typ_polygon"],
             p["podlaorta"], p["lokalita"], p["datum"], p["hlavny_mapovatel"],
             p["druhy_mapovatel"], p["p"], p["plocha_m2"], p["poznamka"],
             p["e0"], p["e1"], p["e2"], p["e3"], p["e1_invaz"], p["e2_invaz"],
             p["e3_invaz"], p["polygon_id_form"], p["zdroj_polygon_id"],
             p["biotopy_suhrn"], p["lat"], p["lng"], p["lat_min"], p["lat_max"],
             p["lng_min"], p["lng_max"], p["prstence"]))
        pfk = cur.lastrowid
        fk_by_fid[p["fid"]] = pfk

        for b in p["biotopy"]:
            perc_cols = [c for c, _ in KVALITA + MANAZMENT + VYHLIADKY]
            bcur = con.execute(
                "INSERT INTO biotop (polygon_fk, poradie, kod2002, nazov2002, "
                "kod2023, nazov2023, pokryv, kvalita, manazment, vyhliadky, "
                "gpkg_fid, %s) VALUES (%s)"
                % (", ".join(perc_cols), ",".join("?" * (11 + len(perc_cols)))),
                [pfk, b["poradie"], b["kod2002"], b["nazov2002"], b["kod2023"],
                 b["nazov2023"], b["pokryv"], b["kvalita"], b["manazment"],
                 b["vyhliadky"], b["gpkg_fid"]]
                + [b["percenta"][c] for c in perc_cols])
            counts["biotop"] += 1
            for o in b["opatrenia"]:
                con.execute(
                    "INSERT INTO opatrenie (biotop_fk, kod, nazov, opis, percento, "
                    "gpkg_fid) VALUES (?,?,?,?,?,?)",
                    (bcur.lastrowid, o["kod"], o["nazov"], o["opis"], o["percento"],
                     o["gpkg_fid"]))
                counts["opatrenie"] += 1

        for d in p["druhy"]:
            con.execute(
                "INSERT INTO druh (polygon_fk, kod, nazov_lat, taxon, pokryvnost, "
                "etaz, pokryvnost_perc, kod_kb, charakteristicky, gpkg_fid) "
                "VALUES (?,?,?,?,?,?,?,?,?,?)",
                (pfk, d["kod"], d["nazov_lat"], d["taxon"], d["pokryvnost"],
                 d["etaz"], d["pokryvnost_perc"], d["kod_kb"], d["charakteristicky"],
                 d["gpkg_fid"]))
            counts["druh"] += 1

        for a in p["aktivity"]:
            con.execute(
                "INSERT INTO aktivita (polygon_fk, kod, nazov, intenzita, "
                "perc_plochy, vplyv, gpkg_fid) VALUES (?,?,?,?,?,?,?)",
                (pfk, a["kod"], a["nazov"], a["intenzita"], a["perc_plochy"],
                 a["vplyv"], a["gpkg_fid"]))
            counts["aktivita"] += 1

        counts["foto"] += insert_fotky(con, site["kod"], pfk, p["fotky"])

    # fotky bez polygónu (web ich môže vypísať zvlášť)
    counts["foto"] += insert_fotky(con, site["kod"], None,
                                   site["fotky_nepriradene"])

    con.executemany("INSERT INTO ciselnik (typ, kod, nazov) VALUES (?,?,?)",
                    site["ciselnik"])
    counts["ciselnik"] = len(site["ciselnik"])

    # odkaz na zdrojový polygón sa dá vyplniť až keď sú vložené všetky riadky
    for p in polygons:
        if p["zdroj_fid"] is not None:
            con.execute("UPDATE polygon SET zdroj_polygon_fk = ? WHERE id = ?",
                        (fk_by_fid.get(p["zdroj_fid"]), fk_by_fid[p["fid"]]))
    return counts


def list_sites(data_dir):
    rows = []
    for path in sorted(glob.glob(os.path.join(
            data_dir, "SKUEV[0-9][0-9][0-9][0-9].sqlite"))):
        try:
            con = sqlite3.connect(path)
            row = con.execute("SELECT kod, nazov, pocet_polygonov, plocha_ha, "
                              "pocet_fotiek, aktualizovane FROM skuev").fetchone()
            con.close()
        except sqlite3.Error as e:
            print("   UPOZORNENIE: %s sa nedá prečítať (%s)." % (path, e))
            continue
        if row:
            rows.append(row)
    if not rows:
        print("Na webe nie je zatiaľ žiadne územie.")
        return
    zobrazene = shown_sites(data_dir)
    print("Územia na webe (%d):" % len(rows))
    for kod, nazov, pocet, plocha, fotiek, akt in rows:
        print("   %s  %-32s %4d polygónov  %10.2f ha  %4d fotiek   %s%s"
              % (kod, nazov or "", pocet or 0, plocha or 0, fotiek or 0,
                 akt or "",
                 "" if kod in zobrazene else "   SKRYTÉ (chýba mapovatelOK)"))


def shown_sites(data_dir):
    """Územia, ktoré index.php zobrazí – majú v datumy.json [mapovatelOK]."""
    try:
        with open(os.path.join(data_dir, DATUMY_JSON), encoding="utf-8") as f:
            datumy = json.load(f)
    except (OSError, ValueError):
        return set()
    return {kod for kod, d in datumy.items() if d.get(SHOW_COLUMN)}


def write_datumy(data_dir):
    """Dátumy odovzdania zo skuev_mapovatel.txt -> <data>\\datumy.json.

    {"SKUEV0035": {"mapovatelOK": "2026-10-08", ...}, ...} – len vyplnené
    dátumy, vo formáte rrrr-mm-dd (zápis dd.mm.rrrr sa prevedie)."""
    datumy = {}
    try:
        with open(MAPOVATEL_FILE, encoding="utf-8-sig", newline="") as f:
            for row in csv.DictReader(f, delimiter="\t"):
                kod = (row.get("skuev") or "").strip()
                if not re.fullmatch(r"SKUEV\d{4}", kod):
                    continue
                for column in DATUMY_COLUMNS:
                    value = (row.get(column) or "").strip()
                    if not value:
                        continue
                    d = parse_datum(value)
                    if d is None:
                        print("   UPOZORNENIE: %s [%s]='%s' nie je dátum – "
                              "vynechané." % (kod, column, value))
                        continue
                    datumy.setdefault(kod, {})[column] = d.isoformat()
    except OSError as e:
        print("   UPOZORNENIE: %s sa nedá prečítať (%s) – dátumy odovzdania "
              "sa neaktualizujú." % (MAPOVATEL_FILE, e))
        return
    path = os.path.join(data_dir, DATUMY_JSON)
    with open(path + ".tmp", "w", encoding="utf-8") as f:
        json.dump(datumy, f, ensure_ascii=False, indent=1, sort_keys=True)
    os.replace(path + ".tmp", path)
    print("Dátumy:   %s (%d území)" % (path, len(datumy)))


def build_site_db(data_dir, site, polygons, predmety, prevod):
    """Zostaví databázu územia pod dočasným názvom a potom ňou nahradí starú,
    aby web nikdy nečítal napoly zapísaný súbor. Vráti počty riadkov."""
    path = site_db_path(data_dir, site["kod"])
    tmp = path + ".tmp"
    if os.path.exists(tmp):
        os.remove(tmp)
    con = open_web_db(tmp)
    try:
        counts = write_site(con, site, polygons)
        counts["predmety"] = write_predmety(con, site["kod"], predmety, prevod)
        con.commit()
    finally:
        con.close()
    try:
        os.replace(tmp, path)
    except OSError as e:
        sys.exit("%s sa nedá prepísať (%s) – nie je otvorená inde? Nová "
                 "databáza ostala v %s." % (path, e, tmp))
    return counts


def update_predmety(data_dir, codes, predmety, prevod):
    """Prepíše len predmety ochrany a prevod kódov v existujúcich databázach
    území (bez čítania cloudu a bez fotiek). Prázdne `codes` = všetky."""
    if not codes:
        codes = [os.path.basename(p)[:-len(".sqlite")] for p in sorted(glob.glob(
            os.path.join(data_dir, "SKUEV[0-9][0-9][0-9][0-9].sqlite")))]
    for code in codes:
        path = site_db_path(data_dir, code)
        if not os.path.isfile(path):
            print("   %s na webe nie je – preskočené." % code)
            continue
        try:
            con = sqlite3.connect(path)
            try:
                pocet = write_predmety(con, code, predmety, prevod)
                con.commit()
            finally:
                con.close()
        except sqlite3.Error as e:
            print("   UPOZORNENIE: %s sa nedá zapísať (%s)." % (path, e))
            continue
        print("   %s  predmety ochrany: %s" % (
            code, "nezapísané" if pocet is None else "%d biotopov" % pocet))


# ----------------------------------------------------------------------
# Hlavná logika
# ----------------------------------------------------------------------

def main():
    parser = argparse.ArgumentParser(
        description="Naplní webové SQLite databázy hotových území.")
    parser.add_argument("skuev", nargs="*", help="kódy území, napr. SKUEV0870 0862")
    parser.add_argument("--data", default=WEB_DATA_DIR,
                        help="priečinok s databázami území (a fotkami)")
    parser.add_argument("--remove", action="store_true",
                        help="územia zadané v argumentoch z webu odstráni")
    parser.add_argument("--list", action="store_true",
                        help="vypíše územia, ktoré sú na webe")
    parser.add_argument("--bez-fotiek", dest="bez_fotiek", action="store_true",
                        help="fotky sa nekopírujú (rýchly refresh údajov)")
    parser.add_argument("--predmety", action="store_true",
                        help="v databázach území na webe (alebo zadaných) len "
                             "prepíše predmety ochrany a prevod kódov Natura 2000")
    args = parser.parse_args()

    data_dir = os.path.abspath(args.data)
    os.makedirs(data_dir, exist_ok=True)
    # fotky idú vedľa databáz: <data>\foto\SKUEV####
    foto_root = os.path.join(data_dir, WEB_FOTO_DIR)
    if os.path.isfile(os.path.join(data_dir, OLD_WEB_DB)):
        print("POZNÁMKA: %s je pôvodná spoločná databáza, web ju už nečíta – "
              "územia z nej pridaj znova a súbor zmaž."
              % os.path.join(data_dir, OLD_WEB_DB))
    write_datumy(data_dir)

    if args.list and not args.skuev:
        list_sites(data_dir)
        return

    codes = []
    for raw in args.skuev:
        code = normalize(raw)
        if code is None:
            sys.exit("Neplatný kód '%s' – očakávam tvar SKUEV#### (napr. SKUEV0870)."
                     % raw)
        if code not in codes:
            codes.append(code)

    if not args.remove:
        predmety, chyba = read_predmety()
        if chyba:
            print("UPOZORNENIE: predmety ochrany sa na web nezapíšu – %s." % chyba)
        prevod, chyba = read_natura_prevod()
        if chyba:
            print("UPOZORNENIE: prevod kódov Natura 2000 sa na web nezapíše – %s."
                  % chyba)
    if args.predmety:
        print("Predmety ochrany v databázach %s:" % data_dir)
        update_predmety(data_dir, codes, predmety, prevod)
        return

    if not codes:
        sys.exit("Zadaj aspoň jedno územie, napr.: python 08_web_hotovo.py SKUEV0870")

    print("Databázy: %s" % data_dir)
    print("Fotky:    %s" % foto_root)

    if args.remove:
        for code in codes:
            removed = remove_site(data_dir, code, foto_root)
            if removed is None:
                print("%s na webe nie je." % code)
            else:
                print("%s odstránené z webu (%d polygónov, %d fotiek)."
                      % ((code,) + removed))
        list_sites(data_dir)
        return

    for code in codes:
        print("\n" + "=" * 64 + "\n%s" % code)
        warnings = []
        folder = find_folder(code)
        site, polygons = read_site(code, folder, warnings, not args.bez_fotiek)
        _, bajtov = export_fotky(site, polygons, foto_root, warnings)
        counts = build_site_db(data_dir, site, polygons, predmety, prevod)
        print("   databáza:   %s" % site_db_path(data_dir, code))
        if code not in shown_sites(data_dir):
            warnings.append("V skuev_mapovatel.txt chýba dátum [%s] – územie "
                            "sa v zozname na webe nezobrazí." % SHOW_COLUMN)
        print("   zdroj:      %s" % site["zdroj"])
        print("   názov:      %s" % (site["nazov"] or "(nenájdený v N2000)"))
        print("   polygóny:   %d (%.2f ha)"
              % (site["pocet_polygonov"], site["plocha_ha"]))
        print("   biotopy:    %d, opatrenia: %d, druhy: %d, aktivity: %d"
              % (counts["biotop"], counts["opatrenie"], counts["druh"],
                 counts["aktivita"]))
        print("   predmety:   %s" % ("nezapísané" if counts["predmety"] is None
                                     else "%d biotopov" % counts["predmety"]))
        if counts["predmety"] == 0:
            warnings.append("Územie nemá v %s žiadne biotopy – predmety ochrany."
                            % os.path.basename(PREDMETY_XLSX))
        if args.bez_fotiek:
            print("   fotky:      preskočené (--bez-fotiek)")
        else:
            nepriradene = len([f for f in site["fotky_nepriradene"]
                               if f.get("subor")])
            print("   fotky:      %d (%.1f MB)%s"
                  % (counts["foto"], bajtov / (1024.0 * 1024.0),
                     ", z toho nepriradených %d" % nepriradene
                     if nepriradene else ""))
        print("   mapovanie:  %s – %s (%s)"
              % (site["datum_od"] or "?", site["datum_do"] or "?",
                 site["mapovatelia"] or "bez mapovateľa"))
        if warnings:
            print("   UPOZORNENIA (%d):" % len(warnings))
            for w in warnings[:20]:
                print("      - %s" % w)
            if len(warnings) > 20:
                print("      ... a ďalších %d" % (len(warnings) - 20))

    print()
    list_sites(data_dir)


if __name__ == "__main__":
    main()
